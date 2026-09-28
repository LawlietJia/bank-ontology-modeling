"""Local structure-preserving intake. Output is evidence, never an approved ontology."""
import csv, json, re, subprocess, tempfile, zipfile
from pathlib import Path
from lxml import etree as ET
from model import file_hash

W='http://schemas.openxmlformats.org/wordprocessingml/2006/main'
NS={'w':W}
def xml(data): return ET.fromstring(data,ET.XMLParser(resolve_entities=False,no_network=True))
def text(node): return ''.join(node.itertext())

def ingest(path,dialect=None,ocr=False,language='chi_sim+eng'):
    path=Path(path)
    result={'schema_version':'1.0','source':{'path':str(path.resolve()),'sha256':file_hash(path),'format':path.suffix.lower()},'blocks':[],'issues':[]}
    def add(anchor,kind,value,disposition='parsed',**extra):
        result['blocks'].append({'anchor':anchor,'kind':kind,'text':value,'disposition':disposition,**extra})
    def issue(code,anchor,message):result['issues'].append(dict(code=code,anchor=anchor,message=message))
    suffix=path.suffix.lower()
    if suffix=='.docx':
        with zipfile.ZipFile(path) as z:
            root=xml(z.read('word/document.xml')); body=root.find('w:body',NS)
            def paragraph(p,anchor):
                accepted=''.join(p.xpath('.//w:t[not(ancestor::w:del)]/text()',namespaces=NS))
                add(anchor,'paragraph',accepted,style=p.xpath('./w:pPr/w:pStyle/@w:val',namespaces=NS))
                for i,rev in enumerate(p.xpath('.//w:ins|.//w:del',namespaces=NS),1):
                    kind='insertion' if rev.tag==f'{{{W}}}ins' else 'deletion'
                    add(anchor+f'/{kind}{i}',kind,''.join(rev.xpath('.//w:t/text()|.//w:delText/text()',namespaces=NS)),'needs_review',parent_anchor=anchor)
                    issue('tracked_change',anchor,'正文含未裁决修订；不能自动采用当前显示文本作为已批准规则')
                for cid in p.xpath('.//w:commentRangeStart/@w:id|.//w:commentReference/@w:id',namespaces=NS):
                    result.setdefault('comment_anchors',{}).setdefault(cid,[]).append(anchor)
                if p.xpath('.//w:drawing|.//w:pict',namespaces=NS): issue('embedded_visual',anchor,'图片或图形须另行视觉复核')
            def table(t,anchor):
                add(anchor,'table_structure','表格',rows=len(t.findall('w:tr',NS)))
                for ri,row in enumerate(t.findall('w:tr',NS),1):
                    for ci,cell in enumerate(row.findall('w:tc',NS),1):
                        for pi,p in enumerate(cell.findall('w:p',NS),1):paragraph(p,f'{anchor}/R{ri}/C{ci}/P{pi}')
                        for ti,nested in enumerate(cell.findall('w:tbl',NS),1):table(nested,f'{anchor}/R{ri}/C{ci}/T{ti}')
                        if cell.xpath('./w:tcPr/w:gridSpan|./w:tcPr/w:vMerge',namespaces=NS):issue('merged_cell',f'{anchor}/R{ri}/C{ci}','合并单元格须结合表头解释')
            np=nt=0
            for el in body:
                if el.tag==f'{{{W}}}p':np+=1;paragraph(el,f'P{np}')
                elif el.tag==f'{{{W}}}tbl':nt+=1;table(el,f'T{nt}')
                elif el.tag!=f'{{{W}}}sectPr':
                    anchor=f'BODY/{len(result["blocks"])+1}'
                    add(anchor,'other_body_content',''.join(el.xpath('.//w:t/text()',namespaces=NS)),'needs_review')
                    issue('unsupported_body_structure',anchor,'内容控件/文本框等特殊结构需要复核')
            for name in z.namelist():
                if name=='word/comments.xml':
                    for c in xml(z.read(name)).findall('w:comment',NS):
                        cid=c.get(f'{{{W}}}id');add('COMMENT/'+cid,'comment',''.join(c.xpath('.//w:t/text()',namespaces=NS)),'needs_review',targets=result.get('comment_anchors',{}).get(cid,[]))
                elif re.match(r'word/(footnotes|endnotes|header\d+|footer\d+)\.xml$',name):
                    for i,p in enumerate(xml(z.read(name)).xpath('//w:p',namespaces=NS),1):paragraph(p,f'{name}/P{i}')
                elif name.startswith(('word/media/','word/embeddings/','word/charts/','word/diagrams/')):
                    issue('embedded_resource',name,'资源保留在原文中，须读取或标明不纳入；未声称自动理解')
    elif suffix=='.pdf':
        import pdfplumber
        with pdfplumber.open(path) as pdf:
            for i,page in enumerate(pdf.pages,1):
                content=page.extract_text(layout=True) or ''
                state='parsed'
                if len(content.strip())<20:
                    state='needs_review'
                    if ocr:
                        with tempfile.TemporaryDirectory() as d:
                            prefix=Path(d)/'page'
                            subprocess.run(['pdftoppm','-f',str(i),'-l',str(i),'-singlefile','-png','-scale-to','2400',str(path),str(prefix)],check=True,capture_output=True,timeout=120)
                            content=subprocess.run(['tesseract',str(prefix)+'.png','stdout','-l',language],check=True,capture_output=True,text=True,timeout=120).stdout
                        issue('ocr_review',f'PAGE{i}','OCR转写需核对原页，尤其数字、否定词与表格')
                    else:issue('scan_or_sparse',f'PAGE{i}','少文本页面可能是扫描件；可用 --ocr 或视觉读取')
                add(f'PAGE{i}','page',content,state)
                for j,table in enumerate(page.extract_tables(),1):add(f'PAGE{i}/T{j}','table',json.dumps(table,ensure_ascii=False))
                if page.images:issue('pdf_visual',f'PAGE{i}','页面含图片；文字提取不能证明图片内容已覆盖')
    elif suffix in {'.xlsx','.xlsm'}:
        import openpyxl
        wb=openpyxl.load_workbook(path,data_only=False); cached=openpyxl.load_workbook(path,data_only=True)
        for sheet in wb:
            add(sheet.title+'!INFO','sheet_structure',json.dumps({'hidden':sheet.sheet_state,'merged':[str(x) for x in sheet.merged_cells.ranges]},ensure_ascii=False))
            for row in sheet:
                for cell in row:
                    if cell.value is None:continue
                    anchor=sheet.title+'!'+cell.coordinate
                    add(anchor,'formula' if cell.data_type=='f' else 'cell',str(cell.value),cached_value=str(cached[sheet.title][cell.coordinate].value) if cell.data_type=='f' else None)
                    if cell.comment:add(anchor+'/COMMENT','comment',cell.comment.text,'needs_review')
        with zipfile.ZipFile(path) as z:
            if any(n.startswith(('xl/media/','xl/charts/','xl/drawings/')) for n in z.namelist()):issue('excel_visual','workbook','工作簿含图表或图形，需视觉复核')
        wb.close();cached.close()
    elif suffix=='.pptx':
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
        prs=Presentation(path)
        def shape_blocks(shape,anchor):
            if shape.shape_type==MSO_SHAPE_TYPE.GROUP:
                for i,s in enumerate(shape.shapes,1):shape_blocks(s,anchor+f'/G{i}')
            if shape.has_text_frame:add(anchor,'slide_text',shape.text)
            if shape.has_table:
                for r,row in enumerate(shape.table.rows,1):
                    for c,cell in enumerate(row.cells,1):add(anchor+f'/R{r}/C{c}','table_cell',cell.text)
            if shape.shape_type==MSO_SHAPE_TYPE.PICTURE or shape.has_chart:issue('slide_visual',anchor,'图片/图表须视觉复核')
        for i,slide in enumerate(prs.slides,1):
            for j,shape in enumerate(slide.shapes,1):shape_blocks(shape,f'SLIDE{i}/S{j}')
            if slide.has_notes_slide:add(f'SLIDE{i}/NOTES','notes',slide.notes_slide.notes_text_frame.text)
    elif suffix=='.sql':
        import sqlglot
        from sqlglot import exp
        result['schema']={'tables':[],'dialect':dialect or 'unspecified'}
        sql=path.read_text(encoding='utf-8-sig')
        try: statements=sqlglot.parse(sql,read=dialect,error_level=sqlglot.ErrorLevel.RAISE)
        except sqlglot.errors.ParseError as exc:
            add('DDL/RAW','ddl',sql,'needs_review');issue('ddl_parse','DDL/RAW',str(exc));statements=[]
        for i,stmt in enumerate(statements,1):
            if stmt is None:continue
            anchor=f'DDL/{i}';add(anchor,'ddl',stmt.sql(dialect=dialect),original_comments=stmt.comments or [])
            if not isinstance(stmt,exp.Create) or stmt.args.get('kind')!='TABLE' or not isinstance(stmt.this,exp.Schema):
                issue('ddl_non_table',anchor,'非CREATE TABLE语句已保留，需单独解释索引、注释、ALTER或其他逻辑');continue
            t={'name':stmt.this.this.sql(dialect=dialect),'anchor':anchor,'columns':[],'primary_key':[],'foreign_keys':[],'constraints':[]}
            for col in stmt.this.expressions:
                if isinstance(col,exp.ColumnDef):
                    constraints=[x.sql(dialect=dialect) for x in col.constraints]
                    t['columns'].append({'name':col.name,'type':col.args['kind'].sql(dialect=dialect),'constraints':constraints,'comments':col.comments or []})
                    if any(isinstance(x.kind,exp.PrimaryKeyColumnConstraint) for x in col.constraints):t['primary_key'].append(col.name)
                    for ref in col.find_all(exp.Reference):
                        target=ref.this
                        t['foreign_keys'].append({'columns':[col.name],'target_table':target.this.sql(dialect=dialect) if isinstance(target,exp.Schema) else str(target),'target_columns':[x.name for x in target.expressions] if isinstance(target,exp.Schema) else []})
                else:t['constraints'].append(col.sql(dialect=dialect))
            for pk in stmt.find_all(exp.PrimaryKey):t['primary_key']=[x.name for x in pk.expressions]
            for fk in stmt.find_all(exp.ForeignKey):
                ref=fk.args.get('reference'); target=ref.this if ref else None
                t['foreign_keys'].append({'columns':[x.name for x in fk.expressions],'target_table':target.this.sql(dialect=dialect) if isinstance(target,exp.Schema) else str(target),'target_columns':[x.name for x in target.expressions] if isinstance(target,exp.Schema) else []})
            result['schema']['tables'].append(t)
        add('DDL/SOURCE','original_ddl',sql)
    elif suffix in {'.png','.jpg','.jpeg','.tif','.tiff'}:
        value=''
        if ocr:value=subprocess.run(['tesseract',str(path),'stdout','-l',language],check=True,capture_output=True,text=True,timeout=120).stdout
        add('IMAGE1','image_ocr',value,'needs_review');issue('image_review','IMAGE1','图像或OCR文字须人工/视觉复核')
    elif suffix in {'.csv','.tsv'}:
        with path.open(encoding='utf-8-sig',newline='') as f:
            for i,row in enumerate(csv.reader(f,delimiter='\t' if suffix=='.tsv' else ','),1):add(f'ROW{i}','row',json.dumps(row,ensure_ascii=False))
    elif suffix in {'.txt','.md','.json','.yaml','.yml','.xml'}:
        for i,line in enumerate(path.read_text(encoding='utf-8-sig').splitlines(),1):add(f'L{i}','text',line)
    else:
        add('FILE','unsupported','', 'needs_review');issue('unsupported_format','FILE','请转换为受支持格式并保留原文件版本')
    result['coverage']={'blocks':len(result['blocks']),'needs_review':sum(b['disposition']!='parsed' for b in result['blocks']),'issues':len(result['issues']),'semantic_extraction':'not_performed'}
    return result

def word_opinions(path):
    parsed=ingest(path)
    return {'source':parsed['source'],'opinions':[b for b in parsed['blocks'] if b['kind'] in {'comment','insertion','deletion'}],'status':'pending_semantic_review','issues':parsed['issues']}
