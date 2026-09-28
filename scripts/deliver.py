"""Business artifacts and lossless standard-workbook editing."""
import copy, json
from pathlib import Path
from model import KINDS, LEGACY_KINDS, STATES, DETAILS, REQUIRED, REF_LIST_FIELDS, digest, validate, write, has_test_evidence

COMMON=[('id','稳定ID'),('name','名称'),('definition','业务定义'),('status','评审状态'),('evidence','证据JSON'),('links','关联ID列表JSON')]
EXTRA={
 'sources':[('id','来源ID'),('path','文件路径'),('kind','来源性质'),('version','来源版本'),('sha256','文件哈希'),('blocks_path','解析件路径')],
 'questions':[('id','问题ID'),('question','问题'),('blocking','是否阻塞'),('status','处理状态'),('affects','影响ID列表JSON'),('resolution','处理决定')],
 'tests':[('id','验收ID'),('question','业务问题'),('expected','预期结果'),('record_ids','关联ID列表JSON'),('result','验证结果'),('evidence','验证记录')],
}
NAMES={'sources':'资料来源','questions':'待确认问题','tests':'业务验收'}
JSON_FIELDS={'evidence','links','affects','record_ids'}

def display(v):
    if v is None:return ''
    if isinstance(v,(dict,list,bool)):return json.dumps(v,ensure_ascii=False)
    return str(v)

def columns(kind,version=2):
    fields=list(dict.fromkeys(['domain','nature','scope']+REQUIRED.get(kind,[])))
    # Supplement fields needed for practical handoff without a fifty-column sheet.
    extra={'relation':['time_basis','condition'],'rule':['time_basis','owner'],'metric':['currency','unknown_handling'],
           'action':['inputs','outputs','state_from','state_to'],'mapping':['source_table','source_field','external_iri','reference_version','alignment'],
           'object':['classification','owner'],'state':['owner_id'],'role':['owner_id','condition']}
    fields=list(dict.fromkeys(fields+extra.get(kind,[])))
    if version>=2:
        fields+=['source_unit_ids']
        if kind in {'concept','object','role'}:
            fields+=['model_level','iri','term_name','labels','synonyms','parent_ids','equivalent_ids','disjoint_ids','instance_of_ids']
        if kind in {'attribute','relation'}:
            fields+=['model_level','iri','term_name','labels','property_type','domain_ids','range_ids','allowed_values','unit']
        if kind=='information_unit':fields+=['about_ids','related_unit_ids','exception_unit_ids','effective_from','effective_to']
        if kind=='constraint':fields+=['time_basis','verification_owner','conformance_basis','interpretation']
        fields=list(dict.fromkeys(fields))
    return COMMON+[("details."+k,DETAILS.get(k,k)) for k in fields]

def get_field(r,key):
    if key.startswith('details.'):return r.get('details',{}).get(key[8:])
    return r.get(key)

def _workbook_fallback(m,path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    wb=Workbook();ws=wb.active;ws.title='使用说明'
    for row in [ ['银行业务本体建模台账'],['项目',m['project']['title']],['版本',m['project']['version']],['状态',m['project']['status']],
        ['编辑约定','修改业务工作表；保持ID、表名和表头。新增填新ID，删除整行会提出删除建议。'],
        ['结构列','证据和关联列使用JSON，证据例 [{"source_id":"S1","anchor":"P1"}]。无法填写时通过对话补充。'],
        ['审批','改为approved不能代替业务批准。回读先形成提案，不自动发布。'],
        ['注意','公式按文字保存；不要把单元格改为可执行公式。隐藏基线用于版本比较，不能编辑。'],
        ['附加字段','未展示的扩展字段保存在基线中；通过对话修改，普通业务修改不会丢失。'] ]:ws.append(row)
    ws.column_dimensions['A'].width=20;ws.column_dimensions['B'].width=95
    for row in ws:
        for cell in row:cell.alignment=Alignment(vertical='top',wrap_text=True)
    ws.row_dimensions[6].height=42
    sheets={}
    for kind,name in KINDS.items():sheets[name]=(columns(kind),[r for r in m['records'] if r['kind']==kind])
    for col,cols in EXTRA.items():sheets[NAMES[col]]=(cols,m[col])
    for name,(cols,records) in sheets.items():
        s=wb.create_sheet(name);s.append([label for _,label in cols]);s.freeze_panes='C2'
        for r in records:
            values=[display(get_field(r,k)) for k,_ in cols]
            if any(len(v)>32767 for v in values):raise ValueError('Excel单元格超过32767字符，须拆分建模内容：'+r['id'])
            s.append(values)
        s.auto_filter.ref=s.dimensions
        for c in s[1]:c.fill=PatternFill('solid',fgColor='153B50');c.font=Font(color='FFFFFF',bold=True);c.alignment=Alignment(wrap_text=True)
        s.row_dimensions[1].height=32
        for j,(key,_) in enumerate(cols,1):s.column_dimensions[get_column_letter(j)].width=20 if j<=2 else 48 if key=='definition' else 32
        for row in s.iter_rows(min_row=2):
            for c in row:
                c.data_type='s';c.number_format='@';c.alignment=Alignment(wrap_text=True,vertical='top')
                if c.row%2==0:c.fill=PatternFill('solid',fgColor='EDF3F6')
            s.row_dimensions[row[0].row].height=90
    meta=wb.create_sheet('_baseline');meta.sheet_state='veryHidden'
    payload=json.dumps(m,ensure_ascii=False,separators=(',',':'))
    meta.append(['bank-ontology-workbook-2',digest(m)])
    for i in range(0,len(payload),20000):meta.append([payload[i:i+20000]])
    Path(path).parent.mkdir(parents=True,exist_ok=True);wb.save(path);wb.close()

def workbook(m,path):
    import os, subprocess, tempfile, zipfile
    from xml.etree import ElementTree as E
    runtime=Path.home()/'.cache/codex-runtimes/codex-primary-runtime/dependencies/node'
    modules=Path(os.environ.get('ONTOLOGY_NODE_MODULES',str(runtime/'node_modules')))
    if not (modules/'@oai/artifact-tool').exists():return _workbook_fallback(m,path)
    node=os.environ.get('ONTOLOGY_NODE',str(runtime/'bin/node'))
    payload=[]
    payload.append({'name':'使用说明','rows':[['银行业务本体建模台账',''],['项目',m['project']['title']],['版本',m['project']['version']],['状态',m['project']['status']],['编辑约定','保持表名、表头与稳定ID。新增填新ID，删除整行会形成删除提案。'],['证据格式','证据使用JSON列表，例如 [{"source_id":"S1","anchor":"P1"}]；可通过对话补充。'],['回读规则','改为approved不等于业务批准；隐藏基线不能编辑。未显示的扩展字段仍会保留。']]})
    for kind,name in KINDS.items():
        cols=columns(kind);payload.append({'name':name,'rows':[[label for _,label in cols]]+[[display(get_field(r,k)) for k,_ in cols] for r in m['records'] if r['kind']==kind]})
    for col,cols in EXTRA.items():payload.append({'name':NAMES[col],'rows':[[label for _,label in cols]]+[[display(r.get(k)) for k,_ in cols] for r in m[col]]})
    body=json.dumps(m,ensure_ascii=False,separators=(',',':'))
    payload.append({'name':'_baseline','rows':[['bank-ontology-workbook-2',digest(m)]]+[[body[i:i+20000],''] for i in range(0,len(body),20000)]})
    for s in payload:
        if any(len(v)>32767 for row in s['rows'] for v in row):raise ValueError('Excel单元格超过32767字符：'+s['name'])
    Path(path).parent.mkdir(parents=True,exist_ok=True)
    with tempfile.TemporaryDirectory() as td:
        src=Path(td)/'sheets.json';write(src,payload)
        generated=Path(td)/'workbook.xlsx'
        subprocess.run([node,str(Path(__file__).with_name('workbook.mjs')),str(src),str(generated),str(modules)],check=True,capture_output=True,text=True,timeout=180)
        with zipfile.ZipFile(generated) as z:parts={n:z.read(n) for n in z.namelist()}
    # The public artifact API does not expose veryHidden in the supplied API guide.
    # Patch only this OOXML metadata flag; spreadsheet values/styles come from artifact-tool.
    root=E.fromstring(parts['xl/workbook.xml']);ns={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    for s in root.findall('s:sheets/s:sheet',ns):
        if s.get('name')=='_baseline':s.set('state','veryHidden')
    parts['xl/workbook.xml']=E.tostring(root,encoding='utf-8',xml_declaration=True)
    with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as z:
        for n,b in parts.items():z.writestr(n,b)

def read_workbook(path):
    from openpyxl import load_workbook
    wb=load_workbook(path,data_only=False)
    try:
        if '_baseline' not in wb:raise ValueError('不是标准导出台账：缺少原基线')
        meta=wb['_baseline']
        formats={'bank-ontology-workbook-1':1,'bank-ontology-workbook-2':2}
        if meta['A1'].value not in formats:raise ValueError('不支持的台账格式')
        version=formats[meta['A1'].value]
        base=json.loads(''.join(str(meta.cell(i,1).value or '') for i in range(2,meta.max_row+1)))
        if digest(base)!=meta['B1'].value:raise ValueError('原基线被修改，无法安全比较')
        edited=copy.deepcopy(base);recs=[]
        kinds=LEGACY_KINDS if version==1 else KINDS
        sheets=[(name,columns(k,version),'records',k) for k,name in kinds.items()]+[(NAMES[col],cols,col,None) for col,cols in EXTRA.items()]
        allowed={'使用说明','_baseline'}|{x[0] for x in sheets}
        if set(wb.sheetnames)!=allowed:raise ValueError('工作表增删或重命名；请保留标准表，通过对话增加视图')
        for name,cols,col,kind in sheets:
            s=wb[name]
            if [c.value for c in s[1]][:len(cols)]!=[v for _,v in cols]:raise ValueError('表头被修改：'+name)
            if any(c.value is not None for row in s.iter_rows(min_col=len(cols)+1) for c in row):raise ValueError('新增列不能静默忽略：'+name)
            original={r['id']:r for r in base[col]};rows=[];seen=set()
            for cells in s.iter_rows(min_row=2,max_col=len(cols)):
                vals=[c.value for c in cells]
                if all(v is None or v=='' for v in vals):continue
                if any(c.data_type=='f' for c in cells):raise ValueError('发现可执行Excel公式，请改成文字：'+name)
                rid=vals[0]
                if not isinstance(rid,str) or not rid.strip():raise ValueError('稳定ID须是非空文本：'+name)
                if rid in seen:raise ValueError('重复ID：'+rid)
                seen.add(rid)
                r=copy.deepcopy(original.get(rid,{'id':rid}))
                if kind:
                    if r.get('kind',kind)!=kind:raise ValueError('不能通过跨表移动改变记录类型：'+rid)
                    r.setdefault('kind',kind);r.setdefault('details',{});r.setdefault('evidence',[]);r.setdefault('links',[]);r.setdefault('status','candidate')
                for (key,_),value in zip(cols,vals):
                    old=get_field(r,key);value='' if value is None else str(value)
                    if value==display(old):continue
                    if key in {'links','affects','record_ids'} or key=='evidence' and col=='records' or key.startswith('details.') and key[8:] in REF_LIST_FIELDS | {'synonyms','representation'}:
                        try:v=json.loads(value or '[]')
                        except ValueError as exc:raise ValueError('JSON格式不正确：'+rid+'/'+key) from exc
                        if not isinstance(v,list):raise ValueError('应为JSON列表：'+rid+'/'+key)
                    elif key in {'details.labels','details.specification'}:
                        try:v=json.loads(value or '{}')
                        except ValueError as exc:raise ValueError('应为JSON对象：'+rid+'/'+key) from exc
                        if not isinstance(v,dict):raise ValueError('应为JSON对象：'+rid+'/'+key)
                    elif key=='blocking':
                        if value.lower() not in {'true','false'}:raise ValueError('是否阻塞填true或false')
                        v=value.lower()=='true'
                    elif isinstance(old,(dict,list)):
                        try:v=json.loads(value)
                        except ValueError as exc:raise ValueError('结构字段须保留JSON格式：'+rid+'/'+key) from exc
                    else:v=value
                    target=r.setdefault('details',{}) if key.startswith('details.') else r
                    target[key.split('.',1)[-1]]=v
                rows.append(r)
            if col=='records':recs+=rows
            else:edited[col]=rows
        if len({r['id'] for r in recs})!=len(recs):raise ValueError('不同业务表存在重复ID')
        # Preserve original ordering; table grouping should not itself be a change.
        order={r['id']:i for i,r in enumerate(base['records'])}
        edited['records']=sorted(recs,key=lambda r:order.get(r['id'],len(order)))
        if any(r['kind'] in {'information_unit','constraint'} for r in recs):edited['schema_version']='1.1'
        return base,edited
    finally:wb.close()

def document(m,path):
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    d=Document();sec=d.sections[0];sec.top_margin=sec.bottom_margin=Inches(.7)
    for name in ['Normal','Title','Heading 1','Heading 2']:
        st=d.styles[name];st.font.name='Arial';st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'PingFang SC');st.font.color.rgb=RGBColor(0,0,0)
        for border in list(st._element.iter(qn('w:pBdr'))):border.getparent().remove(border)
    d.styles['Normal'].font.size=Pt(10);d.styles['Normal'].paragraph_format.space_after=Pt(6)
    d.styles['Title'].font.size=Pt(22)
    d.add_heading(m['project']['title'],0)
    d.add_paragraph('业务模型说明书')
    labels={'status':'处理状态','blocking':'是否阻塞','resolution':'处理决定','affects':'影响记录','expected':'预期结果','result':'验证结果','evidence':'验证依据','record_ids':'关联模型记录'}
    def test_status(t):
        if t.get('result')=='pass':
            return '通过且有验证记录' if has_test_evidence(t) else '声明通过，依据不足'
        return {'not_run':'未执行','fail':'未通过'}.get(str(t.get('result')),'状态无效')
    def readable(value):
        if isinstance(value,bool):return '是' if value else '否'
        if isinstance(value,list):return '；'.join(readable(x) for x in value)
        if isinstance(value,dict):return '；'.join(str(k)+'：'+readable(v) for k,v in value.items())
        return {'open':'待处理','resolved':'已解决','not_run':'未执行','pass':'通过','fail':'未通过','draft':'草案','baseline':'实施基线',**STATES}.get(str(value),display(value))
    enum_labels={
        'model_level':{'class':'类','individual':'实例','property':'属性','statement':'陈述'},
        'property_type':{'data':'数据属性','object':'对象属性'},
        'unit_type':{'requirement':'要求','recommendation':'建议','instruction':'指示','permission':'允许','statement':'陈述','definition':'定义','example':'示例','note':'注释','other':'其他'},
        'statement_nature':{'normative':'规范性内容','fact':'事实','proposed':'拟议内容','example':'示例','modeling_decision':'建模决定','unknown':'待确认'},
        'constraint_type':{'unique':'唯一性','cardinality':'基数','value_range':'数值范围','enumeration':'枚举','date_order':'日期顺序','disjoint':'互斥','reference':'引用','custom':'自定义'},
        'representation':{'text':'文本','table':'表格','figure':'图','formula':'公式'},
        'operator':{'lt':'早于','le':'不晚于'},
        'policy':{'must_exist':'必须存在','allow_external':'允许外部引用','requires_resolution':'需要解析确认'},
    }
    key_labels={'zh':'中文','en':'英文','min':'下限','max':'上限','min_inclusive':'包含下限','max_inclusive':'包含上限',
                'values':'允许值','scope':'适用范围','operator':'顺序关系','policy':'引用策略','expression':'表达式','language':'表达语言'}
    def detail_text(key,value):
        if isinstance(value,bool):return '是' if value else '否'
        if isinstance(value,list):return '；'.join(detail_text(key,x) for x in value) or '未登记'
        if isinstance(value,dict):return '；'.join(key_labels.get(k,k)+'：'+detail_text(k,v) for k,v in value.items()) or '未登记'
        return enum_labels.get(key,{}).get(str(value),display(value))
    def evidence_text(e):
        value=e['source_id']+' @ '+e['anchor']
        if 'field' in e:
            key=e['field'].removeprefix('details.')
            label={'name':'名称','definition':'业务定义'}.get(e['field'],DETAILS.get(key,key))
            value+='〔'+label.removesuffix('JSON')+'〕'
        return value
    d.add_paragraph('项目 '+m['project']['id']+'  版本 '+m['project']['version']+'  状态 '+readable(m['project']['status']))
    d.add_paragraph('本交付用于业务评审和技术交接。'+m['project']['purpose']+'。证据与待确认项随模型保留；草案不代表银行已经批准。')
    d.add_heading('范围与交付约定',1)
    d.add_paragraph('领域：'+readable(m['project']['domains'])+'；应用：'+m['project']['application']+'；业务责任：'+m['project']['owner'])
    for heading,key in [('待确认问题','questions'),('业务验收','tests')]:
        d.add_heading(heading,1)
        if not m[key]:d.add_paragraph('尚未登记；不能据此视为不存在缺口或已完成验收。')
        if key=='tests' and m[key]:
            from collections import Counter
            counts=Counter(test_status(t) for t in m[key])
            d.add_paragraph('已登记验收：'+str(len(m[key]))+' 项；'+'；'.join(k+'：'+str(v) for k,v in counts.items()))
            d.add_paragraph('以上仅汇总已登记的验收状态；验证记录的存在不证明答案正确，也不代替人工业务验收。')
        for item in m[key]:
            d.add_paragraph(item['id']+'  '+item.get('question',''),style='Heading 2')
            for k in (['status','blocking','resolution','affects'] if key=='questions' else ['expected','record_ids','result','evidence']):
                if key=='tests' and k=='result':value=test_status(item)
                elif key=='tests' and k=='record_ids':value=readable(item.get(k)) if item.get(k) else '未关联，待补充适用范围'
                elif k in item:value=readable(item[k])
                else:continue
                d.add_paragraph(labels[k]+'：'+value)
    d.add_heading('业务语义与规则',1)
    for r in m['records']:
        d.add_heading(r['id']+' '+r['name'],2)
        d.add_paragraph(KINDS[r['kind']]+'  '+readable(r['status'])+'。'+r['definition'])
        for k,v in r.get('details',{}).items():
            if v is not None and v!='':d.add_paragraph(DETAILS.get(k,k).removesuffix('JSON')+'：'+detail_text(k,v))
        d.add_paragraph('证据：'+'；'.join(evidence_text(e) for e in r.get('evidence',[])))
    d.add_heading('来源与实施交接',1)
    for s in m['sources']:d.add_paragraph(s['id']+'  '+s.get('path','')+'  版本 '+s.get('version','未登记'))
    d.add_paragraph('结构化模型和Excel台账与本文使用相同ID。数据映射、规则表达、身份及时间口径必须落实；尚无依据的内容以问题为准。动作仅为业务契约，未连接或执行生产系统。Word修改和批注先作为待处理意见，不自动发布。')
    footer=sec.footer.paragraphs[0];footer.add_run(m['project']['id']+'  '+m['project']['version']+'  |  ')
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
    d.save(path)

def diagram(m,path):
    from xml.etree import ElementTree as E
    outer=E.Element('mxfile',host='bank-ontology-modeling'); pages={}
    objects={r['id']:r for r in m['records'] if r['kind'] in {'object','concept','role','event','state','decision','action'}}
    for r in objects.values():pages.setdefault(r['details'].get('domain','共享概念'),[]).append(r)
    if not pages:pages={'模型说明':[]}
    for name,items in pages.items():
        dg=E.SubElement(outer,'diagram',name=name);graph=E.SubElement(dg,'mxGraphModel',adaptiveColors='auto');root=E.SubElement(graph,'root')
        E.SubElement(root,'mxCell',id='0');E.SubElement(root,'mxCell',id='1',parent='0')
        local={r['id'] for r in items}; relations=[r for r in m['records'] if r['kind']=='relation' and (r['details'].get('subject') in local or r['details'].get('object') in local)]
        needed={r['details'][k] for r in relations for k in ['subject','object']} - local
        nodes=items+[objects[k] for k in sorted(needed) if k in objects]
        for i,r in enumerate(nodes):
            fill='#edf3f6' if r['id'] in local else '#fff2cc'
            cell=E.SubElement(root,'mxCell',id=r['id'],value=r['id']+' '+r['name'],vertex='1',parent='1',style='rounded=1;whiteSpace=wrap;html=0;fillColor='+fill+';strokeColor=#153b50;fontSize=14;')
            E.SubElement(cell,'mxGeometry',x=str(40+(i%3)*300),y=str(40+(i//3)*180),width='220',height='75',attrib={'as':'geometry'})
        for i,r in enumerate(relations):
            a=r['details']['subject'];b=r['details']['object']
            if a not in objects or b not in objects:continue
            cell=E.SubElement(root,'mxCell',id='edge-'+r['id'],value=r['name']+' '+r['details'].get('cardinality',''),edge='1',parent='1',source=a,target=b,style='edgeStyle=orthogonalEdgeStyle;rounded=0;html=0;endArrow=open;fontSize=11;labelBackgroundColor=#ffffff;')
            E.SubElement(cell,'mxGeometry',relative='1',attrib={'as':'geometry'})
    E.ElementTree(outer).write(path,encoding='utf-8',xml_declaration=True)

def export(m,out,base_dir=None):
    import shutil
    import tempfile
    errors=[x for x in validate(m,base_dir) if x['severity']=='error']
    if errors:raise ValueError('模型校验未通过：'+str(errors))
    out=Path(out).absolute()
    if out.exists() or out.is_symlink():raise ValueError('输出目录已存在，请使用尚不存在的新版本目录')
    out.parent.mkdir(parents=True,exist_ok=True)
    targets=['model.json','业务建模台账.xlsx','业务模型说明书.docx','业务模型.drawio','validation.json']
    m=copy.deepcopy(m)
    with tempfile.TemporaryDirectory(prefix='.'+out.name+'-',dir=out.parent) as temp:
        stage=Path(temp)
        if base_dir:
            for index,s in enumerate(m['sources'],1):
                if s.get('blocks_path'):
                    source=Path(s['blocks_path']);source=source if source.is_absolute() else Path(base_dir)/source
                    dest=stage/'evidence'/f'source-{index}.json';dest.parent.mkdir(exist_ok=True)
                    shutil.copyfile(source,dest);s['blocks_path']=dest.relative_to(stage).as_posix()
        write(stage/'model.json',m);workbook(m,stage/'业务建模台账.xlsx');document(m,stage/'业务模型说明书.docx');diagram(m,stage/'业务模型.drawio')
        issues=validate(m,stage)
        if any(x['severity']=='error' for x in issues):raise ValueError('交付证据校验未通过：'+str(issues))
        write(stage/'validation.json',{'model_hash':digest(m),'issues':issues,'semantic_review':'not_certified_by_script',
            'validation_scope':'model_structure_evidence_and_constraint_specifications',
            'policy_execution':'not_run','shacl_validation':'not_run','gbt48000_conformance':'not_assessed'})
        # Reserve the final path exclusively after generation; never replace an existing directory.
        out.mkdir()
        try:
            stage.rename(out)
        except BaseException:
            try:out.rmdir()  # Only our empty reservation; preserve any concurrent writer's files.
            except OSError:pass
            raise
    return {'directory':str(out),'files':targets,'version':m['project']['version']}
